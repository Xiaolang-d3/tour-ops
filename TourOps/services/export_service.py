"""导出服务 - PDF/Excel 生成"""
from io import BytesIO
from datetime import datetime, date, timedelta
from decimal import Decimal
from typing import List, Dict, Any, Optional
from sqlalchemy.orm import Session

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm, mm
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    HRFlowable, KeepTogether, PageBreak
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.graphics.shapes import Drawing, Rect, String
from reportlab.graphics import renderPDF

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

from TourOps.models.trip import Trip
from TourOps.models.activity import Activity, ActivityType


# ======================== 颜色常量 ========================
BRAND_BLUE = colors.HexColor('#3366CC')
BRAND_DARK = colors.HexColor('#2C3E50')
BRAND_GRAY = colors.HexColor('#7F8C8D')
LIGHT_BG = colors.HexColor('#F8F9FA')
BORDER_COLOR = colors.HexColor('#DEE2E6')
DAY_COLORS = [
    colors.HexColor('#3366CC'),
    colors.HexColor('#2E86C1'),
    colors.HexColor('#1ABC9C'),
    colors.HexColor('#27AE60'),
    colors.HexColor('#F39C12'),
    colors.HexColor('#E74C3C'),
    colors.HexColor('#9B59B6'),
]


class ExportService:
    def __init__(self, db: Session):
        self.db = db
        self._font_name = self._register_fonts()

    def _register_fonts(self) -> str:
        """注册中文字体，返回可用的字体名"""
        font_paths = [
            ('SimHei', 'simhei.ttf'),
            ('SimHei', 'SimHei.ttf'),
            ('MicrosoftYaHei', 'msyh.ttc'),
            ('MicrosoftYaHei', 'msyh.ttf'),
            ('SimSun', 'simsun.ttc'),
        ]
        for font_name, font_file in font_paths:
            try:
                pdfmetrics.registerFont(TTFont(font_name, font_file))
                return font_name
            except Exception:
                continue
        # 尝试系统字体路径
        import os
        win_font_dir = r'C:\Windows\Fonts'
        for font_name, font_file in font_paths:
            try:
                full_path = os.path.join(win_font_dir, font_file)
                pdfmetrics.registerFont(TTFont(font_name, full_path))
                return font_name
            except Exception:
                continue
        return 'Helvetica'

    def _create_styles(self) -> dict:
        """创建所有 PDF 样式"""
        fn = self._font_name
        return {
            'title': ParagraphStyle(
                'PDFTitle', fontName=fn, fontSize=20, leading=28,
                textColor=BRAND_DARK, spaceAfter=2,
            ),
            'subtitle': ParagraphStyle(
                'PDFSubtitle', fontName=fn, fontSize=11, leading=16,
                textColor=BRAND_GRAY,
            ),
            'date_right': ParagraphStyle(
                'DateRight', fontName=fn, fontSize=9, leading=14,
                textColor=BRAND_GRAY, alignment=TA_RIGHT,
            ),
            'section_title': ParagraphStyle(
                'SectionTitle', fontName=fn, fontSize=15, leading=22,
                textColor=BRAND_BLUE, spaceBefore=18, spaceAfter=10,
            ),
            'day_title': ParagraphStyle(
                'DayTitle', fontName=fn, fontSize=13, leading=20,
                textColor=BRAND_DARK, spaceAfter=6,
            ),
            'time_text': ParagraphStyle(
                'TimeText', fontName=fn, fontSize=10, leading=16,
                textColor=BRAND_DARK,
            ),
            'activity_text': ParagraphStyle(
                'ActivityText', fontName=fn, fontSize=10, leading=16,
                textColor=colors.HexColor('#4A4A4A'),
            ),
            'normal': ParagraphStyle(
                'PDFNormal', fontName=fn, fontSize=10, leading=15,
                textColor=colors.HexColor('#4A4A4A'),
            ),
            'note_title': ParagraphStyle(
                'NoteTitle', fontName=fn, fontSize=12, leading=18,
                textColor=BRAND_BLUE, spaceAfter=6,
            ),
            'note_text': ParagraphStyle(
                'NoteText', fontName=fn, fontSize=9, leading=14,
                textColor=colors.HexColor('#5A5A5A'),
            ),
            'toc_text': ParagraphStyle(
                'TOCText', fontName=fn, fontSize=10, leading=20,
                textColor=colors.HexColor('#4A4A4A'),
            ),
            'footer_text': ParagraphStyle(
                'FooterText', fontName=fn, fontSize=8, leading=12,
                textColor=BRAND_GRAY, alignment=TA_CENTER,
            ),
        }

    # ======================== 数据获取 ========================

    def get_trip_with_activities(self, trip_id: int) -> tuple:
        """获取行程及其活动"""
        trip = self.db.query(Trip).filter(Trip.id == trip_id).first()
        if not trip:
            return None, []
        activities = self.db.query(Activity).filter(
            Activity.trip_id == trip_id
        ).order_by(Activity.start_time).all()
        return trip, activities

    # ======================== PDF 生成 ========================

    def generate_pdf(self, trip_id: int) -> BytesIO:
        """生成专业版 PDF 行程单"""
        trip, activities = self.get_trip_with_activities(trip_id)
        if not trip:
            raise ValueError("行程不存在")

        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=2 * cm,
            leftMargin=2 * cm,
            topMargin=1.8 * cm,
            bottomMargin=1.8 * cm,
        )

        styles = self._create_styles()
        story = []

        # ---- 1. 页眉 Header ----
        story.extend(self._build_header(trip, styles))

        # ---- 2. 分割线 ----
        story.append(Spacer(1, 0.4 * cm))
        story.append(HRFlowable(
            width="100%", thickness=1, color=BRAND_BLUE,
            spaceAfter=0.6 * cm,
        ))

        # ---- 3. 行程概览信息 ----
        story.extend(self._build_overview(trip, styles))

        # ---- 4. 目录 Table of Contents ----
        activities_by_date = self._group_activities_by_date(activities, trip)
        story.extend(self._build_toc(activities_by_date, trip, styles))

        # ---- 5. 每日行程 Daily Schedule ----
        story.append(Spacer(1, 0.4 * cm))
        story.append(Paragraph("每日行程", styles['section_title']))
        story.append(Spacer(1, 0.2 * cm))

        for idx, (date_key, day_info) in enumerate(activities_by_date.items()):
            day_block = self._build_day_section(
                idx, day_info, trip, styles
            )
            if len(day_info['activities']) <= 6:
                story.append(KeepTogether(day_block))
            else:
                story.extend(day_block)
            story.append(Spacer(1, 0.3 * cm))

        # ---- 6. 费用汇总 ----
        total_cost = sum(act.cost or Decimal(0) for act in activities)
        if total_cost > 0:
            story.extend(self._build_cost_summary(activities, total_cost, styles))

        # ---- 7. 底部：注意事项 + 联系信息 ----
        story.append(Spacer(1, 0.6 * cm))
        story.extend(self._build_footer_notes(trip, styles))

        # ---- 8. 页脚生成时间 ----
        story.append(Spacer(1, 0.8 * cm))
        story.append(HRFlowable(
            width="100%", thickness=0.5, color=BORDER_COLOR,
            spaceAfter=0.3 * cm,
        ))
        story.append(Paragraph(
            f"由 TourOps 行程编排系统生成 · {datetime.now().strftime('%Y-%m-%d %H:%M')}",
            styles['footer_text'],
        ))

        doc.build(story)
        buffer.seek(0)
        return buffer

    # ============================================================
    #  PDF 各区块构建方法
    # ============================================================

    def _build_header(self, trip: Trip, styles: dict) -> list:
        """构建页眉：Logo + 标题 + 日期"""
        fn = self._font_name
        elements = []

        # Logo 方块
        logo_drawing = Drawing(44, 44)
        logo_drawing.add(Rect(0, 0, 44, 44, rx=6, ry=6,
                              fillColor=BRAND_BLUE, strokeColor=None))
        logo_drawing.add(String(10, 14, "TP",
                                fontName='Helvetica-Bold', fontSize=18,
                                fillColor=colors.white))

        # 标题区域
        title_para = Paragraph("行程计划书", styles['title'])
        days = (trip.end_date - trip.start_date).days + 1
        subtitle_para = Paragraph(
            f"{trip.name} · {days}天行程",
            styles['subtitle'],
        )

        # 日期
        date_para = Paragraph(
            f"生成日期: {datetime.now().strftime('%Y年%m月%d日')}",
            styles['date_right'],
        )

        # 用 Table 组装三列布局
        header_table = Table(
            [[logo_drawing, [title_para, subtitle_para], date_para]],
            colWidths=[1.8 * cm, 10.5 * cm, 4.5 * cm],
        )
        header_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))
        elements.append(header_table)
        return elements

    def _build_overview(self, trip: Trip, styles: dict) -> list:
        """构建行程概览信息条"""
        fn = self._font_name
        elements = []
        days = (trip.end_date - trip.start_date).days + 1

        info_items = [
            f"📅 {self._format_date_cn(trip.start_date)} — {self._format_date_cn(trip.end_date)}",
            f"⏱ 共 {days} 天",
            f"👥 {trip.guest_count} 人",
        ]
        if trip.budget:
            info_items.append(f"💰 预算 ¥{trip.budget:,.0f}")

        info_text = "　　".join(info_items)
        info_style = ParagraphStyle(
            'OverviewInfo', parent=styles['normal'],
            fontSize=9, textColor=BRAND_GRAY,
        )
        elements.append(Paragraph(info_text, info_style))
        elements.append(Spacer(1, 0.4 * cm))
        return elements

    def _build_toc(self, activities_by_date: dict, trip: Trip, styles: dict) -> list:
        """构建目录"""
        fn = self._font_name
        elements = []
        elements.append(Paragraph("目　录", styles['section_title']))
        elements.append(Spacer(1, 0.2 * cm))

        toc_data = []
        for idx, (date_key, day_info) in enumerate(activities_by_date.items()):
            day_num = idx + 1
            day_label = day_info.get('theme', f"第{day_num}天")
            date_str = day_info['date'].strftime('%m月%d日') if day_info.get('date') else ''
            toc_data.append([
                Paragraph(
                    f"第 {day_num} 天: {day_label}",
                    styles['toc_text'],
                ),
                Paragraph(date_str, ParagraphStyle(
                    'TOCDate', parent=styles['toc_text'],
                    alignment=TA_RIGHT, textColor=BRAND_GRAY,
                )),
            ])

        # 注意事项条目
        toc_data.append([
            Paragraph("注意事项 & 联系方式", styles['toc_text']),
            Paragraph("", styles['toc_text']),
        ])

        toc_table = Table(toc_data, colWidths=[12 * cm, 4.8 * cm])
        toc_style = [
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (0, -1), 12),
            ('RIGHTPADDING', (-1, 0), (-1, -1), 12),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('LINEBELOW', (0, 0), (-1, -2), 0.5, colors.HexColor('#E0E0E0')),
        ]
        toc_table.setStyle(TableStyle(toc_style))
        elements.append(toc_table)
        return elements

    def _build_day_section(self, idx: int, day_info: dict,
                           trip: Trip, styles: dict) -> list:
        """构建单日行程区块（带左侧彩色边框）"""
        fn = self._font_name
        elements = []
        day_num = idx + 1
        day_color = DAY_COLORS[idx % len(DAY_COLORS)]
        act_date = day_info.get('date')
        theme = day_info.get('theme', '')
        day_acts = day_info.get('activities', [])

        # 构建日期标题
        date_display = act_date.strftime('%Y年%m月%d日') if act_date else ''
        title_text = f"<b>第 {day_num} 天</b>"
        if theme:
            title_text += f" - {theme}"
        if date_display:
            title_text += f"  <font color='#7F8C8D' size='10'>({date_display})</font>"

        day_title = Paragraph(title_text, styles['day_title'])

        # 构建活动列表
        act_rows = []
        for act in day_acts:
            time_str = act.start_time.strftime("%H:%M") if act.start_time else ""
            act_name = act.name or ""
            location_str = ""
            if act.location:
                location_str = f"  <font color='#999999' size='8'>📍{act.location}</font>"

            time_para = Paragraph(
                f"<b>{time_str}</b>",
                styles['time_text'],
            )
            detail_para = Paragraph(
                f"{act_name}{location_str}",
                styles['activity_text'],
            )
            act_rows.append([time_para, detail_para])

        if not act_rows:
            act_rows.append([
                Paragraph("", styles['normal']),
                Paragraph("<font color='#999'>暂无活动安排</font>", styles['normal']),
            ])

        # 活动表格
        act_table = Table(act_rows, colWidths=[2.2 * cm, 13 * cm])
        act_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (0, -1), 0),
            ('LEFTPADDING', (1, 0), (1, -1), 10),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))

        # 将标题和活动列表合并为内容
        inner_content = [day_title, Spacer(1, 0.15 * cm), act_table]

        # 用外层 Table 实现左侧彩色边框效果
        # 第1列: 彩色竖条（4pt宽）, 第2列: 内容
        content_cell = inner_content
        border_table = Table(
            [["", content_cell]],
            colWidths=[4, 15.6 * cm],
        )
        border_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, 0), day_color),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (0, 0), 0),
            ('RIGHTPADDING', (0, 0), (0, 0), 0),
            ('LEFTPADDING', (1, 0), (1, 0), 12),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            # 浅灰底色
            ('BACKGROUND', (1, 0), (1, 0), colors.HexColor('#FAFBFC')),
            ('ROUNDEDCORNERS', [0, 4, 4, 0]),
        ]))

        elements.append(border_table)
        return elements

    def _build_cost_summary(self, activities: list, total_cost: Decimal,
                            styles: dict) -> list:
        """构建费用汇总区块"""
        fn = self._font_name
        elements = []
        elements.append(Spacer(1, 0.3 * cm))
        elements.append(Paragraph("费用汇总", styles['section_title']))

        cost_by_type = self._calculate_cost_by_type(activities)
        cost_data = []

        # 表头
        header_style = ParagraphStyle(
            'CostHeader', fontName=fn, fontSize=9, leading=14,
            textColor=colors.white,
        )
        header_style_r = ParagraphStyle(
            'CostHeaderR', parent=header_style, alignment=TA_RIGHT,
        )
        cost_data.append([
            Paragraph("类　型", header_style),
            Paragraph("金　额", header_style_r),
        ])

        row_style = ParagraphStyle(
            'CostRow', fontName=fn, fontSize=10, leading=16,
            textColor=BRAND_DARK,
        )
        row_style_r = ParagraphStyle(
            'CostRowR', parent=row_style, alignment=TA_RIGHT,
        )

        for type_name, cost in cost_by_type.items():
            cost_data.append([
                Paragraph(type_name, row_style),
                Paragraph(f"¥{cost:,.0f}", row_style_r),
            ])

        # 合计行
        total_style = ParagraphStyle(
            'CostTotal', fontName=fn, fontSize=11, leading=18,
            textColor=BRAND_DARK,
        )
        total_style_r = ParagraphStyle(
            'CostTotalR', parent=total_style, alignment=TA_RIGHT,
        )
        cost_data.append([
            Paragraph("<b>合　计</b>", total_style),
            Paragraph(f"<b>¥{total_cost:,.0f}</b>", total_style_r),
        ])

        cost_table = Table(cost_data, colWidths=[8 * cm, 8.8 * cm])
        table_style = [
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BACKGROUND', (0, 0), (-1, 0), BRAND_BLUE),
            ('LEFTPADDING', (0, 0), (-1, -1), 12),
            ('RIGHTPADDING', (0, 0), (-1, -1), 12),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('LINEBELOW', (0, 1), (-1, -3), 0.5, BORDER_COLOR),
            ('LINEABOVE', (0, -1), (-1, -1), 1, BRAND_BLUE),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#EBF5FB')),
        ]
        cost_table.setStyle(TableStyle(table_style))
        elements.append(cost_table)
        return elements

    def _build_footer_notes(self, trip: Trip, styles: dict) -> list:
        """构建底部注意事项和联系信息两栏"""
        fn = self._font_name
        elements = []

        # 分割线
        elements.append(HRFlowable(
            width="100%", thickness=0.5, color=BORDER_COLOR,
            spaceAfter=0.4 * cm,
        ))

        # 左栏：注意事项
        notes_title = Paragraph("<b>注意事项</b>", styles['note_title'])
        notes_items = [
            "• 请提前 15 分钟到达集合地点",
            "• 请携带有效身份证件",
            "• 如有特殊饮食需求请提前告知",
            "• 行程可能因天气等原因调整",
            "• 请保管好个人贵重物品",
        ]

        # 如果行程有特殊要求，添加到注意事项
        if trip.special_requirements:
            if isinstance(trip.special_requirements, dict):
                notes_text = trip.special_requirements.get('notes', '')
                if notes_text:
                    notes_items.append(f"• {notes_text}")
            elif isinstance(trip.special_requirements, str):
                notes_items.append(f"• {trip.special_requirements}")

        notes_paras = [notes_title]
        for note in notes_items:
            notes_paras.append(Paragraph(note, styles['note_text']))

        # 右栏：联系信息
        contact_title = Paragraph("<b>联系信息</b>", styles['note_title'])
        contact_items = [
            "运营方: TourOps 行程编排系统",
            "客服邮箱: support@tourops.com",
            f"行程编号: #{trip.id:06d}",
            f"分享码: {trip.share_code[:8] if trip.share_code else 'N/A'}",
        ]
        contact_paras = [contact_title]
        for item in contact_items:
            contact_paras.append(Paragraph(item, styles['note_text']))

        footer_table = Table(
            [[notes_paras, contact_paras]],
            colWidths=[9 * cm, 7.8 * cm],
        )
        footer_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))
        elements.append(footer_table)
        return elements

    # ============================================================
    #  工具方法
    # ============================================================

    def _group_activities_by_date(self, activities: List[Activity],
                                  trip: Trip) -> Dict[str, dict]:
        """按日期分组活动，附带日期和主题"""
        grouped: Dict[str, dict] = {}

        # 先为行程的每一天创建空分组
        days = (trip.end_date - trip.start_date).days + 1
        for i in range(days):
            d = trip.start_date + timedelta(days=i)
            key = d.isoformat()
            grouped[key] = {
                'date': d,
                'theme': self._guess_day_theme(i + 1, []),
                'activities': [],
            }

        # 填充活动
        for act in activities:
            date_key = act.start_time.strftime("%Y-%m-%d")
            if date_key not in grouped:
                d = act.start_time.date()
                grouped[date_key] = {
                    'date': d,
                    'theme': '',
                    'activities': [],
                }
            grouped[date_key]['activities'].append(act)

        # 根据活动推断每日主题
        for key, info in grouped.items():
            if info['activities']:
                info['theme'] = self._guess_day_theme_from_activities(info['activities'])

        return dict(sorted(grouped.items()))

    def _guess_day_theme_from_activities(self, activities: List[Activity]) -> str:
        """根据活动内容推断当日主题"""
        type_counts = {}
        for act in activities:
            t = act.type
            type_counts[t] = type_counts.get(t, 0) + 1

        # 找出最多的非餐饮/交通类型
        main_types = {k: v for k, v in type_counts.items()
                      if k not in (ActivityType.MEAL, ActivityType.TRANSPORT)}
        if main_types:
            dominant = max(main_types, key=main_types.get)
            theme_map = {
                ActivityType.ATTRACTION: "景点游览",
                ActivityType.HOTEL: "住宿休整",
                ActivityType.FREE: "自由活动",
            }
            return theme_map.get(dominant, "综合行程")

        # 全是餐饮/交通
        if ActivityType.MEAL in type_counts:
            return "美食体验"
        return "行程安排"

    def _guess_day_theme(self, day_num: int, activities: list) -> str:
        """默认主题"""
        return f"第{day_num}天行程"

    def _format_date_cn(self, d) -> str:
        """格式化日期为中文"""
        if isinstance(d, (date, datetime)):
            return d.strftime('%m月%d日')
        return str(d)

    def _get_activity_type_name(self, activity_type: ActivityType) -> str:
        """获取活动类型中文名"""
        type_names = {
            ActivityType.TRANSPORT: "交通",
            ActivityType.ATTRACTION: "景点",
            ActivityType.MEAL: "餐饮",
            ActivityType.HOTEL: "住宿",
            ActivityType.FREE: "自由活动",
        }
        return type_names.get(activity_type, str(activity_type))

    def _calculate_cost_by_type(self, activities: List[Activity]) -> Dict[str, Decimal]:
        """按类型统计费用"""
        cost_by_type = {}
        for act in activities:
            if act.cost:
                type_name = self._get_activity_type_name(act.type)
                cost_by_type[type_name] = cost_by_type.get(type_name, Decimal(0)) + act.cost
        return cost_by_type

    # ======================== Excel 生成 ========================

    def generate_excel(self, trip_id: int) -> BytesIO:
        """生成 Excel 明细表"""
        trip, activities = self.get_trip_with_activities(trip_id)
        if not trip:
            raise ValueError("行程不存在")

        wb = Workbook()
        ws = wb.active
        ws.title = "行程明细"

        # 样式定义
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="3366CC", end_color="3366CC", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='DEE2E6'),
            right=Side(style='thin', color='DEE2E6'),
            top=Side(style='thin', color='DEE2E6'),
            bottom=Side(style='thin', color='DEE2E6'),
        )
        title_font = Font(bold=True, size=18, color='2C3E50')
        subtitle_font = Font(size=11, color='7F8C8D')
        info_label_font = Font(bold=True, size=10, color='2C3E50')
        info_value_font = Font(size=10, color='4A4A4A')

        # 标题行
        ws.merge_cells('A1:G1')
        ws['A1'] = trip.name
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 36

        # 副标题
        days = (trip.end_date - trip.start_date).days + 1
        ws.merge_cells('A2:G2')
        ws['A2'] = f"行程计划书 · {days}天行程 · 生成于 {datetime.now().strftime('%Y-%m-%d')}"
        ws['A2'].font = subtitle_font
        ws.row_dimensions[2].height = 22

        # 空行
        ws.row_dimensions[3].height = 8

        # 基本信息
        info_data = [
            ('出行日期', f"{trip.start_date} 至 {trip.end_date}"),
            ('出行人数', f"{trip.guest_count} 人"),
            ('行程天数', f"{days} 天"),
        ]
        if trip.budget:
            info_data.append(('预　算', f"¥{trip.budget:,.0f}"))

        for i, (label, value) in enumerate(info_data):
            row = 4 + i
            ws.cell(row=row, column=1, value=label).font = info_label_font
            ws.cell(row=row, column=2, value=value).font = info_value_font
            ws.row_dimensions[row].height = 22

        # 空行
        data_start = 4 + len(info_data) + 1
        ws.row_dimensions[data_start - 1].height = 8

        # 活动明细表头
        headers = ["日期", "时间", "类型", "活动名称", "地点", "费用", "备注"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=data_start, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        ws.row_dimensions[data_start].height = 28

        # 活动数据
        row = data_start + 1
        alt_fill = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
        for i, act in enumerate(activities):
            is_alt = i % 2 == 1
            cells = [
                ws.cell(row=row, column=1, value=act.start_time.strftime("%Y-%m-%d")),
                ws.cell(row=row, column=2, value=act.start_time.strftime("%H:%M")),
                ws.cell(row=row, column=3, value=self._get_activity_type_name(act.type)),
                ws.cell(row=row, column=4, value=act.name),
                ws.cell(row=row, column=5, value=act.location or ""),
                ws.cell(row=row, column=6, value=float(act.cost) if act.cost else 0),
                ws.cell(row=row, column=7, value=act.notes or ""),
            ]
            for cell in cells:
                cell.border = thin_border
                cell.font = Font(size=10, color='4A4A4A')
                if is_alt:
                    cell.fill = alt_fill
            # 费用列右对齐
            cells[5].alignment = Alignment(horizontal='right')
            cells[5].number_format = '#,##0'
            ws.row_dimensions[row].height = 24
            row += 1

        # 合计行
        total_cost = sum(float(act.cost or 0) for act in activities)
        ws.cell(row=row, column=5, value="合计").font = Font(bold=True, size=11, color='2C3E50')
        ws.cell(row=row, column=5).alignment = Alignment(horizontal='right')
        total_cell = ws.cell(row=row, column=6, value=total_cost)
        total_cell.font = Font(bold=True, size=11, color='3366CC')
        total_cell.number_format = '#,##0'
        total_cell.alignment = Alignment(horizontal='right')
        ws.row_dimensions[row].height = 28

        # 调整列宽
        ws.column_dimensions['A'].width = 14
        ws.column_dimensions['B'].width = 8
        ws.column_dimensions['C'].width = 10
        ws.column_dimensions['D'].width = 28
        ws.column_dimensions['E'].width = 22
        ws.column_dimensions['F'].width = 12
        ws.column_dimensions['G'].width = 22

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
